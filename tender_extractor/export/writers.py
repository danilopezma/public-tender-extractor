"""Deliverables: concurrencia.json, incidencias.json and concurrencia.xlsx."""

import json
from pathlib import Path
from typing import List, Optional

from openpyxl import Workbook
from openpyxl.styles import Font

from tender_extractor.models import Issue, Offer, Tender

CONCURRENCIA_FILE = "concurrencia.json"
INCIDENCIAS_FILE = "incidencias.json"
XLSX_FILE = "concurrencia.xlsx"

# (header, value from an offer, number format). Amounts and points are written as
# Decimal, so the cell holds exactly the figure read from the document.
_MONEY = "#,##0.00"
_COLUMNS = [
    ("Lote", lambda o: o.lot, None),
    ("Orden", lambda o: o.order, None),
    ("Licitador", lambda o: o.bidder.name, None),
    ("NIF", lambda o: o.bidder.tax_id, None),
    ("Precio unitario licitación (sin IVA)", lambda o: o.tender_unit_price, _MONEY),
    ("Precio unitario ofertado (sin IVA)", lambda o: o.offered_unit_price, _MONEY),
    ("Total licitación (sin IVA)", lambda o: o.tender_total, _MONEY),
    ("Total ofertado (sin IVA)", lambda o: o.offered_total, _MONEY),
    ("Puntos criterio precio", lambda o: o.price_points, _MONEY),
    ("Puntos criterios cualitativos", lambda o: o.quality_points, _MONEY),
    ("Total puntuación", lambda o: o.total_points, _MONEY),
    ("Orden de prelación", lambda o: o.ranking, None),
    ("Adjudicatario propuesto", lambda o: "Sí" if o.proposed_awardee else "No", None),
    ("Baja anormal o temeraria", lambda o: "Sí" if o.abnormally_low else "No", None),
    ("Posiciones baja anormal", lambda o: ", ".join(o.abnormally_low_positions) or None, None),
]


def write_xlsx(path: Path, tender: Optional[Tender], offers: List[Offer]) -> None:
    """Write the Excel file: one row per lot, order and bidder, plus the expediente.

    Args:
        path: File to write.
        tender: Expediente, or None if it could not be read.
        offers: Offers, one per row.
    """
    workbook = Workbook()

    # 1. Concurrency sheet: one row per offer ---------------------------------
    sheet = workbook.active
    sheet.title = "Concurrencia"
    sheet.append([header for header, _, _ in _COLUMNS])
    for offer in offers:
        sheet.append([get(offer) for _, get, _ in _COLUMNS])
    for column, (_, _, number_format) in enumerate(_COLUMNS, start=1):
        sheet.cell(row=1, column=column).font = Font(bold=True)
        if number_format:
            for row in range(2, sheet.max_row + 1):
                sheet.cell(row=row, column=column).number_format = number_format

    # 2. Expediente sheet -----------------------------------------------------
    info = workbook.create_sheet("Expediente")
    for label, value in [
        ("Número de expediente", tender.file_number if tender else None),
        ("Objeto", tender.subject if tender else None),
        ("Órgano de contratación", tender.contracting_authority if tender else None),
        ("Fecha de la mesa", tender.meeting_date if tender else None),
    ]:
        info.append([label, value])
    workbook.save(path)


def write_concurrencia(path: Path, tender: Optional[Tender], offers: List[Offer]) -> None:
    """Write concurrencia.json: the expediente and one object per offer.

    Amounts and scores are written as strings, so no figure goes through a float.

    Args:
        path: File to write.
        tender: Expediente, or None if it could not be read.
        offers: Offers.
    """
    path.write_text(json.dumps({
        "tender": tender.model_dump(mode="json") if tender else None,
        "offers": [offer.model_dump(mode="json") for offer in offers],
    }, indent=2, ensure_ascii=False), encoding="utf-8")


def write_incidencias(path: Path, issues: List[Issue]) -> None:
    """Write incidencias.json: the problems found.

    Args:
        path: File to write.
        issues: Issues to write.
    """
    path.write_text(json.dumps([issue.model_dump(mode="json") for issue in issues], indent=2, ensure_ascii=False), encoding="utf-8")
