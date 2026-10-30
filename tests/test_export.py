import json

from openpyxl import load_workbook

from tender_extractor.export.writers import write_concurrencia, write_xlsx
from tender_extractor.models import Bidder, Offer, Tender


def test_xlsx_has_one_row_per_offer_and_json_keeps_amounts_as_strings(tmp_path):
    offer = Offer(lot="1", order="1.10", bidder=Bidder(name="ALFA S.A."), offered_unit_price="12789.50")
    write_xlsx(tmp_path / "c.xlsx", Tender(file_number="AM 1/2025"), [offer])
    write_concurrencia(tmp_path / "c.json", Tender(file_number="AM 1/2025"), [offer])

    sheet = load_workbook(tmp_path / "c.xlsx")["Concurrencia"]
    assert sheet.max_row == 2 and sheet["B2"].value == "1.10"
    assert json.loads((tmp_path / "c.json").read_text())["offers"][0]["offered_unit_price"] == "12789.50"
